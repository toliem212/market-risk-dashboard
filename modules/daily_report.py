"""Báo cáo giám sát rủi ro thị trường hằng ngày."""

from __future__ import annotations

from io import BytesIO

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


from modules.app_config import CANDIDATE_NAME


STATUS_NORMAL = "Bình thường"
STATUS_WATCH = "Theo dõi"
STATUS_ALERT = "Cảnh báo"


STATUS_RANK = {
    STATUS_NORMAL: 0,
    STATUS_WATCH: 1,
    STATUS_ALERT: 2,
}


def _latest_row(frame: pd.DataFrame) -> pd.Series | None:
    """Lấy quan sát mới nhất theo cột date."""
    if frame.empty or "date" not in frame.columns:
        return None

    valid = frame.copy()
    valid["date"] = pd.to_datetime(valid["date"], errors="coerce")
    valid = valid.dropna(subset=["date"]).sort_values("date")

    if valid.empty:
        return None

    return valid.iloc[-1]


def _safe_float(value: object) -> float:
    """Chuyển giá trị về float, trả NaN nếu không hợp lệ."""
    try:
        if pd.isna(value):
            return np.nan
        return float(value)
    except (TypeError, ValueError):
        return np.nan


def _append_snapshot_row(
    rows: list[dict[str, object]],
    date: object,
    group: str,
    indicator: str,
    value: object,
    unit: str,
    status: str = STATUS_NORMAL,
) -> None:
    """Thêm một chỉ báo vào bảng snapshot thị trường."""
    rows.append(
        {
            "Ngày dữ liệu": pd.to_datetime(date, errors="coerce"),
            "Nhóm giám sát": group,
            "Chỉ báo": indicator,
            "Giá trị": _safe_float(value),
            "Đơn vị": unit,
            "Trạng thái": status,
        }
    )


def build_market_snapshot(
    fx_data: pd.DataFrame,
    interbank_data: pd.DataFrame,
    deposit_data: pd.DataFrame,
    funding_pressure_data: pd.DataFrame,
    alert_snapshot: pd.DataFrame | None = None,
    bond_monitor_data: pd.DataFrame | None = None,
) -> pd.DataFrame:
    """Tạo bảng chỉ báo thị trường tại quan sát mới nhất của từng nguồn."""
    rows: list[dict[str, object]] = []

    status_lookup: dict[tuple[str, str], str] = {}
    if alert_snapshot is not None and not alert_snapshot.empty:
        for _, row in alert_snapshot.iterrows():
            key = (str(row.get("group", "")), str(row.get("indicator", "")))
            status_lookup[key] = str(row.get("status", STATUS_NORMAL))

    fx = _latest_row(fx_data)
    if fx is not None:
        fx_date = fx["date"]
        _append_snapshot_row(
            rows,
            fx_date,
            "Ngoại hối",
            "USD/VND đóng cửa",
            fx.get("close"),
            "VND/USD",
        )
        _append_snapshot_row(
            rows,
            fx_date,
            "Ngoại hối",
            "Biến động tỷ giá trong ngày",
            fx.get("daily_return_pct"),
            "%",
            status_lookup.get(("Ngoại hối", "Biến động tỷ giá trong ngày"), STATUS_NORMAL),
        )
        _append_snapshot_row(
            rows,
            fx_date,
            "Ngoại hối",
            "Biên độ tỷ giá trong ngày",
            fx.get("intraday_range_pct"),
            "%",
            status_lookup.get(("Ngoại hối", "Biên độ tỷ giá trong ngày"), STATUS_NORMAL),
        )
        _append_snapshot_row(
            rows,
            fx_date,
            "Ngoại hối",
            "Độ biến động tỷ giá 20 ngày",
            fx.get("volatility_20d_pct"),
            "%/năm",
            status_lookup.get(("Ngoại hối", "Độ biến động tỷ giá 20 ngày"), STATUS_NORMAL),
        )

    interbank = _latest_row(interbank_data)
    if interbank is not None:
        ib_date = interbank["date"]
        _append_snapshot_row(
            rows,
            ib_date,
            "Liên ngân hàng",
            "Lãi suất O/N",
            interbank.get("rate_on"),
            "%",
            status_lookup.get(("Liên ngân hàng", "Mặt bằng lãi suất O/N"), STATUS_NORMAL),
        )
        _append_snapshot_row(
            rows,
            ib_date,
            "Liên ngân hàng",
            "Biến động lãi suất O/N trong ngày",
            interbank.get("on_change_bps"),
            "bps",
            status_lookup.get(
                ("Liên ngân hàng", "Biến động lãi suất O/N trong ngày"),
                STATUS_NORMAL,
            ),
        )
        _append_snapshot_row(
            rows,
            ib_date,
            "Liên ngân hàng",
            "Lãi suất 1 tháng",
            interbank.get("rate_1m"),
            "%",
        )
        _append_snapshot_row(
            rows,
            ib_date,
            "Liên ngân hàng",
            "Lãi suất 3 tháng",
            interbank.get("rate_3m"),
            "%",
        )
        spread_bps = _safe_float(interbank.get("spread_3m_on")) * 100
        _append_snapshot_row(
            rows,
            ib_date,
            "Liên ngân hàng",
            "Chênh lệch 3 tháng - O/N",
            spread_bps,
            "bps",
        )
        turnover_column = (
            "total_turnover_monitoring"
            if "total_turnover_monitoring" in interbank_data.columns
            else "total_turnover"
        )
        turnover_data = interbank_data.dropna(subset=[turnover_column])
        turnover_row = _latest_row(turnover_data)
        if turnover_row is not None:
            _append_snapshot_row(
                rows,
                turnover_row["date"],
                "Liên ngân hàng",
                "Tổng doanh số giao dịch",
                turnover_row.get(turnover_column),
                "tỷ đồng",
                status_lookup.get(
                    ("Liên ngân hàng", "Doanh số giao dịch liên ngân hàng"),
                    STATUS_NORMAL,
                ),
            )

    deposit = _latest_row(deposit_data)
    if deposit is not None:
        dep_date = deposit["date"]
        _append_snapshot_row(
            rows,
            dep_date,
            "Lãi suất huy động",
            "Lãi suất 1-3 tháng",
            deposit.get("deposit_1_3m"),
            "%",
        )
        _append_snapshot_row(
            rows,
            dep_date,
            "Lãi suất huy động",
            "Lãi suất 6-9 tháng",
            deposit.get("deposit_6_9m"),
            "%",
        )
        _append_snapshot_row(
            rows,
            dep_date,
            "Lãi suất huy động",
            "Lãi suất 12 tháng",
            deposit.get("deposit_12m"),
            "%",
            status_lookup.get(
                ("Lãi suất huy động", "Mặt bằng lãi suất huy động 12 tháng"),
                STATUS_NORMAL,
            ),
        )

        if len(deposit_data) > 20:
            change_20 = (
                pd.to_numeric(deposit_data["deposit_12m"], errors="coerce").diff(20).iloc[-1]
                * 100
            )
        else:
            change_20 = np.nan
        _append_snapshot_row(
            rows,
            dep_date,
            "Lãi suất huy động",
            "Thay đổi lãi suất 12 tháng trong 20 quan sát",
            change_20,
            "bps",
            status_lookup.get(
                (
                    "Lãi suất huy động",
                    "Điều chỉnh lãi suất 12 tháng trong 20 quan sát",
                ),
                STATUS_NORMAL,
            ),
        )

        spread_deposit = (
            _safe_float(deposit.get("deposit_12m"))
            - _safe_float(deposit.get("deposit_1_3m"))
        ) * 100
        _append_snapshot_row(
            rows,
            dep_date,
            "Lãi suất huy động",
            "Chênh lệch 12 tháng - 1-3 tháng",
            spread_deposit,
            "bps",
            status_lookup.get(
                ("Lãi suất huy động", "Chênh lệch 12 tháng - 1-3 tháng"),
                STATUS_NORMAL,
            ),
        )

    funding_valid = funding_pressure_data.dropna(
        subset=["funding_pressure_index"]
    ) if not funding_pressure_data.empty else pd.DataFrame()
    if not funding_valid.empty:
        funding = funding_valid.sort_values("date").iloc[-1]
        _append_snapshot_row(
            rows,
            funding["date"],
            "Áp lực nguồn vốn",
            "Chỉ số áp lực nguồn vốn",
            funding.get("funding_pressure_index"),
            "điểm",
            str(funding.get("status", STATUS_NORMAL)),
        )

    if bond_monitor_data is not None and not bond_monitor_data.empty:
        bonds = bond_monitor_data.copy()
        bonds["date"] = pd.to_datetime(bonds["date"], errors="coerce")
        latest_date = bonds["date"].max()
        latest_curve = bonds[bonds["date"] == latest_date].sort_values("tenor_years")

        for _, row in latest_curve.iterrows():
            tenor = _safe_float(row.get("tenor_years"))
            if pd.isna(tenor):
                continue
            _append_snapshot_row(
                rows,
                row.get("date"),
                "Trái phiếu Chính phủ",
                f"Lợi suất {tenor:g} năm",
                row.get("yield_pct"),
                "%",
                str(row.get("status", STATUS_NORMAL)),
            )
            _append_snapshot_row(
                rows,
                row.get("date"),
                "Trái phiếu Chính phủ",
                f"Biến động lợi suất {tenor:g} năm",
                row.get("yield_change_bps"),
                "bps",
                str(row.get("status", STATUS_NORMAL)),
            )

    if not rows:
        return pd.DataFrame(
            columns=[
                "Ngày dữ liệu",
                "Nhóm giám sát",
                "Chỉ báo",
                "Giá trị",
                "Đơn vị",
                "Trạng thái",
            ]
        )

    result = pd.DataFrame(rows)
    result["Ngày dữ liệu"] = pd.to_datetime(result["Ngày dữ liệu"], errors="coerce")
    result["severity_rank"] = result["Trạng thái"].map(STATUS_RANK).fillna(0)
    result = result.sort_values(
        ["severity_rank", "Nhóm giám sát", "Chỉ báo"],
        ascending=[False, True, True],
    )
    return result.drop(columns="severity_rank").reset_index(drop=True)


def build_source_status(
    fx_data: pd.DataFrame,
    interbank_data: pd.DataFrame,
    deposit_data: pd.DataFrame,
    funding_pressure_data: pd.DataFrame,
    bond_monitor_data: pd.DataFrame | None = None,
) -> pd.DataFrame:
    """Tổng hợp phạm vi ngày và số quan sát của từng nguồn dữ liệu."""
    frames: list[tuple[str, pd.DataFrame]] = [
        ("USD/VND", fx_data),
        ("Liên ngân hàng", interbank_data),
        ("Lãi suất huy động", deposit_data),
        ("Áp lực nguồn vốn", funding_pressure_data),
    ]

    if bond_monitor_data is not None and not bond_monitor_data.empty:
        frames.append(("Trái phiếu Chính phủ", bond_monitor_data))

    rows: list[dict[str, object]] = []
    for name, frame in frames:
        if frame.empty or "date" not in frame.columns:
            continue

        dates = pd.to_datetime(frame["date"], errors="coerce").dropna()
        rows.append(
            {
                "Nguồn dữ liệu": name,
                "Số quan sát": int(len(frame)),
                "Từ ngày": dates.min() if not dates.empty else pd.NaT,
                "Đến ngày": dates.max() if not dates.empty else pd.NaT,
            }
        )

    return pd.DataFrame(rows)


def build_exception_report(
    alert_snapshot: pd.DataFrame,
    quality_exceptions: pd.DataFrame,
) -> pd.DataFrame:
    """Gom ngoại lệ thị trường và ngoại lệ dữ liệu vào một bảng vận hành."""
    rows: list[dict[str, object]] = []

    if not alert_snapshot.empty:
        alerts = alert_snapshot[
            alert_snapshot["status"].isin([STATUS_WATCH, STATUS_ALERT])
        ].copy()

        for _, row in alerts.iterrows():
            rows.append(
                {
                    "Ngày": pd.to_datetime(row.get("date"), errors="coerce"),
                    "Nguồn": "Giám sát thị trường",
                    "Nhóm": row.get("group", ""),
                    "Nội dung": row.get("indicator", ""),
                    "Mức độ": row.get("status", ""),
                    "Giá trị": row.get("value", np.nan),
                    "Đơn vị": row.get("value_unit", ""),
                    "Ngưỡng theo dõi": row.get("watch_threshold", np.nan),
                    "Ngưỡng cảnh báo": row.get("alert_threshold", np.nan),
                    "Trạng thái xử lý": "",
                    "Ghi chú xử lý": "",
                }
            )

    if not quality_exceptions.empty:
        for _, row in quality_exceptions.iterrows():
            rows.append(
                {
                    "Ngày": pd.to_datetime(row.get("Ngày"), errors="coerce"),
                    "Nguồn": "Kiểm soát dữ liệu",
                    "Nhóm": row.get("Bộ dữ liệu", ""),
                    "Nội dung": row.get("Loại ngoại lệ", ""),
                    "Mức độ": row.get("Mức độ", ""),
                    "Giá trị": np.nan,
                    "Đơn vị": "",
                    "Ngưỡng theo dõi": np.nan,
                    "Ngưỡng cảnh báo": np.nan,
                    "Trạng thái xử lý": "",
                    "Ghi chú xử lý": row.get("Chi tiết", ""),
                }
            )

    columns = [
        "Ngày",
        "Nguồn",
        "Nhóm",
        "Nội dung",
        "Mức độ",
        "Giá trị",
        "Đơn vị",
        "Ngưỡng theo dõi",
        "Ngưỡng cảnh báo",
        "Trạng thái xử lý",
        "Ghi chú xử lý",
    ]

    if not rows:
        return pd.DataFrame(columns=columns)

    result = pd.DataFrame(rows, columns=columns)
    result["severity_rank"] = result["Mức độ"].map(STATUS_RANK).fillna(0)
    result = result.sort_values(
        ["severity_rank", "Ngày"],
        ascending=[False, False],
    )
    return result.drop(columns="severity_rank").reset_index(drop=True)


def summarize_daily_report(
    alert_snapshot: pd.DataFrame,
    quality_report: pd.DataFrame,
    quality_exceptions: pd.DataFrame,
    source_status: pd.DataFrame,
) -> dict[str, object]:
    """Tạo các chỉ tiêu tổng hợp cho báo cáo ngày."""
    if alert_snapshot.empty:
        watch_count = 0
        alert_count = 0
    else:
        watch_count = int(alert_snapshot["status"].eq(STATUS_WATCH).sum())
        alert_count = int(alert_snapshot["status"].eq(STATUS_ALERT).sum())

    flagged_checks = 0
    if not quality_report.empty and "Số trường hợp" in quality_report.columns:
        flagged_checks = int(quality_report["Số trường hợp"].gt(0).sum())

    latest_date = pd.NaT
    if not source_status.empty:
        latest_date = pd.to_datetime(source_status["Đến ngày"], errors="coerce").max()

    if alert_count > 0:
        overall_status = STATUS_ALERT
    elif watch_count > 0:
        overall_status = STATUS_WATCH
    else:
        overall_status = STATUS_NORMAL

    return {
        "latest_date": latest_date,
        "overall_status": overall_status,
        "open_alert_count": watch_count + alert_count,
        "watch_count": watch_count,
        "alert_count": alert_count,
        "flagged_check_count": flagged_checks,
        "data_exception_count": int(len(quality_exceptions)),
    }


def _prepare_for_excel(frame: pd.DataFrame) -> pd.DataFrame:
    """Chuẩn hóa kiểu dữ liệu trước khi ghi Excel."""
    result = frame.copy()

    for column in result.columns:
        if pd.api.types.is_datetime64_any_dtype(result[column]):
            result[column] = result[column].dt.tz_localize(None)

    return result


def _write_dataframe(writer: pd.ExcelWriter, frame: pd.DataFrame, sheet_name: str) -> None:
    """Ghi một DataFrame vào workbook."""
    _prepare_for_excel(frame).to_excel(writer, sheet_name=sheet_name, index=False)


def _style_workbook(writer: pd.ExcelWriter) -> None:
    """Áp dụng định dạng thống nhất cho workbook."""
    workbook = writer.book
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    alert_fill = PatternFill("solid", fgColor="F4CCCC")
    watch_fill = PatternFill("solid", fgColor="FFF2CC")
    normal_fill = PatternFill("solid", fgColor="D9EAD3")

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, pd.Timestamp):
                    cell.number_format = "dd/mm/yyyy"

                if cell.value == STATUS_ALERT:
                    cell.fill = alert_fill
                elif cell.value == STATUS_WATCH:
                    cell.fill = watch_fill
                elif cell.value == STATUS_NORMAL:
                    cell.fill = normal_fill

        for column_index, column_cells in enumerate(worksheet.columns, start=1):
            max_length = 0
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(value))

            width = min(max(max_length + 2, 11), 45)
            worksheet.column_dimensions[get_column_letter(column_index)].width = width

        worksheet.row_dimensions[1].height = 24


def export_daily_report_excel(
    report_summary: dict[str, object],
    market_snapshot: pd.DataFrame,
    alert_snapshot: pd.DataFrame,
    alert_history: pd.DataFrame,
    group_status: pd.DataFrame,
    quality_report: pd.DataFrame,
    quality_exceptions: pd.DataFrame,
    source_status: pd.DataFrame,
    reconciliation: pd.DataFrame,
    exception_report: pd.DataFrame,
    bond_monitor_data: pd.DataFrame | None = None,
) -> bytes:
    """Xuất báo cáo giám sát ngày ra Excel."""
    summary_table = pd.DataFrame(
        [
            {"Chỉ tiêu": "Ứng viên", "Giá trị": CANDIDATE_NAME},
            {"Chỉ tiêu": "Ngày dữ liệu mới nhất", "Giá trị": report_summary.get("latest_date")},
            {"Chỉ tiêu": "Trạng thái tổng hợp", "Giá trị": report_summary.get("overall_status")},
            {"Chỉ tiêu": "Tín hiệu đang mở", "Giá trị": report_summary.get("open_alert_count", 0)},
            {"Chỉ tiêu": "Tín hiệu Theo dõi", "Giá trị": report_summary.get("watch_count", 0)},
            {"Chỉ tiêu": "Tín hiệu Cảnh báo", "Giá trị": report_summary.get("alert_count", 0)},
            {"Chỉ tiêu": "Phép kiểm tra dữ liệu có ngoại lệ", "Giá trị": report_summary.get("flagged_check_count", 0)},
            {"Chỉ tiêu": "Ngoại lệ dữ liệu có chi tiết", "Giá trị": report_summary.get("data_exception_count", 0)},
        ]
    )

    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        _write_dataframe(writer, summary_table, "Tong_quan")
        _write_dataframe(writer, market_snapshot, "Thi_truong")
        _write_dataframe(writer, group_status, "Trang_thai_nhom")
        _write_dataframe(writer, alert_snapshot, "Canh_bao_hien_tai")
        _write_dataframe(writer, alert_history, "Lich_su_canh_bao")
        _write_dataframe(writer, exception_report, "Exception_Report")
        _write_dataframe(writer, quality_report, "Kiem_soat_du_lieu")
        _write_dataframe(writer, quality_exceptions, "Ngoai_le_du_lieu")
        _write_dataframe(writer, source_status, "Nguon_du_lieu")
        _write_dataframe(writer, reconciliation, "Doi_chieu_du_lieu")

        if bond_monitor_data is not None and not bond_monitor_data.empty:
            bond_export = bond_monitor_data.copy()
            selected_columns = [
                column
                for column in [
                    "date",
                    "tenor_years",
                    "yield_pct",
                    "yield_change_bps",
                    "watch_threshold_bps",
                    "alert_threshold_bps",
                    "status",
                ]
                if column in bond_export.columns
            ]
            _write_dataframe(writer, bond_export[selected_columns], "Duong_cong_loi_suat")

        _style_workbook(writer)

    return buffer.getvalue()


def build_report_filename(latest_date: object) -> str:
    """Tạo tên file báo cáo theo ngày dữ liệu mới nhất."""
    date = pd.to_datetime(latest_date, errors="coerce")
    suffix = date.strftime("%Y%m%d") if pd.notna(date) else "latest"
    return f"To_Thanh_Liem_Market_Risk_Daily_Report_{suffix}.xlsx"
